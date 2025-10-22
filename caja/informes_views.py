from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.db.models import Sum, Q, Count
from django.utils import timezone
from datetime import datetime, date
from decimal import Decimal
import calendar
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO

from .models import CajaDiaria, Recreo, EventoEspecial, PagoProveedor


@login_required
def informes_menu(request):
    """Vista principal del menú de informes"""
    now = timezone.now()
    context = {
        'mes_actual': now.month,
        'anio_actual': now.year,
    }
    return render(request, 'informes/menu.html', context)


@login_required 
def informe_ventas(request):
    """Vista para mostrar informe de ventas con filtros"""
    # Obtener parámetros de filtro
    mes = request.GET.get('mes')
    anio = request.GET.get('anio', timezone.now().year)
    turno = request.GET.get('turno', '')
    nivel = request.GET.get('nivel', '')
    recreo_num = request.GET.get('recreo', '')
    
    # Construir filtros
    filtros = Q()
    
    if mes:
        filtros &= Q(fecha__month=mes)
    if anio:
        filtros &= Q(fecha__year=anio)
    if turno:
        filtros &= Q(turno=turno)
    if nivel:
        filtros &= Q(nivel=nivel)
    
    # Obtener cajas filtradas
    cajas = CajaDiaria.objects.filter(filtros).order_by('-fecha', 'turno', 'nivel')
    
    # Calcular totales
    total_saldo_inicial = cajas.aggregate(Sum('saldo_inicial'))['saldo_inicial__sum'] or 0
    
    # Obtener recreos filtrados
    recreos_filtros = Q(caja__in=cajas)
    if recreo_num:
        recreos_filtros &= Q(numero=recreo_num)
    
    recreos = Recreo.objects.filter(recreos_filtros).select_related('caja')
    total_recreos = recreos.aggregate(Sum('monto'))['monto__sum'] or 0
    
    # Obtener eventos especiales
    eventos = EventoEspecial.objects.filter(caja__in=cajas).select_related('caja')
    total_eventos = eventos.aggregate(Sum('monto'))['monto__sum'] or 0
    
    # Obtener pagos a proveedores
    pagos = PagoProveedor.objects.filter(caja__in=cajas).select_related('caja', 'proveedor')
    total_pagos = pagos.aggregate(Sum('monto'))['monto__sum'] or 0
    
    # Agrupar datos por caja
    datos_cajas = []
    for caja in cajas:
        recreos_caja = recreos.filter(caja=caja)
        eventos_caja = eventos.filter(caja=caja)
        pagos_caja = pagos.filter(caja=caja)
        
        total_ingresos = (recreos_caja.aggregate(Sum('monto'))['monto__sum'] or 0) + \
                        (eventos_caja.aggregate(Sum('monto'))['monto__sum'] or 0)
        total_egresos = pagos_caja.aggregate(Sum('monto'))['monto__sum'] or 0
        
        datos_cajas.append({
            'caja': caja,
            'recreos': recreos_caja,
            'eventos': eventos_caja, 
            'pagos': pagos_caja,
            'total_ingresos': total_ingresos,
            'total_egresos': total_egresos,
            'saldo_final': caja.saldo_inicial + total_ingresos - total_egresos
        })
    
    # Totales generales
    total_ingresos_general = total_recreos + total_eventos
    saldo_final_general = total_saldo_inicial + total_ingresos_general - total_pagos
    
    context = {
        'datos_cajas': datos_cajas,
        'total_saldo_inicial': total_saldo_inicial,
        'total_recreos': total_recreos,
        'total_eventos': total_eventos,
        'total_ingresos_general': total_ingresos_general,
        'total_pagos': total_pagos,
        'saldo_final_general': saldo_final_general,
        'mes': mes,
        'anio': anio,
        'turno': turno,
        'nivel': nivel,
        'recreo_num': recreo_num,
        'meses': [(i, calendar.month_name[i]) for i in range(1, 13)],
        'turnos': CajaDiaria.TURNOS,
        'niveles': CajaDiaria.NIVELES,
        'recreos_disponibles': range(1, 5)  # Recreos 1-4
    }
    
    return render(request, 'informes/ventas.html', context)


@login_required
def exportar_ventas_excel(request):
    """Exportar informe de ventas a Excel"""
    # Obtener los mismos filtros que la vista de informe
    mes = request.GET.get('mes')
    anio = request.GET.get('anio', timezone.now().year)
    turno = request.GET.get('turno', '')
    nivel = request.GET.get('nivel', '')
    recreo_num = request.GET.get('recreo', '')
    
    # Aplicar mismos filtros
    filtros = Q()
    
    if mes:
        filtros &= Q(fecha__month=mes)
    if anio:
        filtros &= Q(fecha__year=anio)
    if turno:
        filtros &= Q(turno=turno)
    if nivel:
        filtros &= Q(nivel=nivel)
    
    cajas = CajaDiaria.objects.filter(filtros).order_by('-fecha', 'turno', 'nivel')
    
    # Obtener datos relacionados
    recreos_filtros = Q(caja__in=cajas)
    if recreo_num:
        recreos_filtros &= Q(numero=recreo_num)
    
    recreos = Recreo.objects.filter(recreos_filtros).select_related('caja')
    eventos = EventoEspecial.objects.filter(caja__in=cajas).select_related('caja')
    pagos = PagoProveedor.objects.filter(caja__in=cajas).select_related('caja', 'proveedor')
    
    # Crear libro de Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Informe de Ventas"
    
    # Estilos
    titulo_font = Font(name='Arial', size=14, bold=True)
    header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Título del reporte
    ws.merge_cells('A1:H1')
    ws['A1'] = 'INFORME DE VENTAS - SISTEMA KIOSKO'
    ws['A1'].font = titulo_font
    ws['A1'].alignment = Alignment(horizontal='center')
    
    # Información de filtros
    row = 3
    filtros_info = []
    if mes:
        mes_nombre = calendar.month_name[int(mes)]
        filtros_info.append(f"Mes: {mes_nombre}")
    if anio:
        filtros_info.append(f"Año: {anio}")
    if turno:
        turno_nombre = dict(CajaDiaria.TURNOS)[turno]
        filtros_info.append(f"Turno: {turno_nombre}")
    if nivel:
        nivel_nombre = dict(CajaDiaria.NIVELES)[nivel]
        filtros_info.append(f"Nivel: {nivel_nombre}")
    if recreo_num:
        filtros_info.append(f"Recreo: {recreo_num}")
    
    if filtros_info:
        ws.merge_cells(f'A{row}:H{row}')
        ws[f'A{row}'] = f"Filtros aplicados: {' | '.join(filtros_info)}"
        row += 2
    
    # Encabezados
    encabezados = ['Fecha', 'Turno', 'Nivel', 'Saldo Inicial', 'Total Recreos', 'Total Eventos', 'Total Egresos', 'Saldo Final']
    
    for col, encabezado in enumerate(encabezados, 1):
        cell = ws.cell(row=row, column=col)
        cell.value = encabezado
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center')
    
    row += 1
    
    # Datos de cajas
    total_saldo_inicial = 0
    total_recreos_sum = 0
    total_eventos_sum = 0
    total_egresos_sum = 0
    total_saldo_final = 0
    
    for caja in cajas:
        recreos_caja = recreos.filter(caja=caja)
        eventos_caja = eventos.filter(caja=caja)
        pagos_caja = pagos.filter(caja=caja)
        
        total_recreos_caja = recreos_caja.aggregate(Sum('monto'))['monto__sum'] or 0
        total_eventos_caja = eventos_caja.aggregate(Sum('monto'))['monto__sum'] or 0
        total_pagos_caja = pagos_caja.aggregate(Sum('monto'))['monto__sum'] or 0
        saldo_final_caja = caja.saldo_inicial + total_recreos_caja + total_eventos_caja - total_pagos_caja
        
        # Datos de la fila
        datos = [
            caja.fecha.strftime('%d/%m/%Y'),
            caja.get_turno_display(),
            caja.get_nivel_display(),
            float(caja.saldo_inicial),
            float(total_recreos_caja),
            float(total_eventos_caja),
            float(total_pagos_caja),
            float(saldo_final_caja)
        ]
        
        for col, valor in enumerate(datos, 1):
            cell = ws.cell(row=row, column=col)
            cell.value = valor
            cell.border = border
            if col > 3:  # Columnas de montos
                cell.number_format = '#,##0.00'
        
        # Acumular totales
        total_saldo_inicial += caja.saldo_inicial
        total_recreos_sum += total_recreos_caja
        total_eventos_sum += total_eventos_caja
        total_egresos_sum += total_pagos_caja
        total_saldo_final += saldo_final_caja
        
        row += 1
    
    # Fila de totales
    row += 1
    totales = ['TOTALES', '', '', float(total_saldo_inicial), float(total_recreos_sum), 
              float(total_eventos_sum), float(total_egresos_sum), float(total_saldo_final)]
    
    for col, valor in enumerate(totales, 1):
        cell = ws.cell(row=row, column=col)
        cell.value = valor
        cell.font = Font(bold=True)
        cell.border = border
        if col > 3 and valor:  # Columnas de montos
            cell.number_format = '#,##0.00'
    
    # Ajustar ancho de columnas
    for col in range(1, 9):
        ws.column_dimensions[get_column_letter(col)].width = 15
    
    # Preparar respuesta HTTP
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    
    # Generar nombre de archivo
    filename = f"informe_ventas"
    if mes:
        filename += f"_{calendar.month_name[int(mes)]}"
    if anio:
        filename += f"_{anio}"
    if turno:
        filename += f"_{dict(CajaDiaria.TURNOS)[turno]}"
    if nivel:
        filename += f"_{dict(CajaDiaria.NIVELES)[nivel]}"
    
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    
    # Guardar y enviar
    wb.save(response)
    return response